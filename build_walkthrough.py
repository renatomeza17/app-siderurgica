from pathlib import Path
import openpyxl

root=Path.cwd()
base=root/'app/src/main/java/pe/edu/unmsm/siderurgica'
parts=[]
verified=[]
def say(s): parts.append(s.strip()+'\n')
def code(f,a,b=None):
    p=base/f if f.endswith('.kt') else root/f
    lines=p.read_text(encoding='utf-8-sig').splitlines()
    b=len(lines) if b is None else b
    assert 1<=a<=b<=len(lines),(f,a,b)
    snippet='\n'.join(lines[a-1:b])
    parts.append(f'*En pantalla: `{p.relative_to(root).as_posix()}`, líneas {a}–{b}.*\n\n```{"kotlin" if f.endswith(".kt") else "text"}\n{snippet}\n```\n')
    verified.append((p,a,b,snippet))
def pick(f,needle,count=1):
    lines=(base/f).read_text(encoding='utf-8-sig').splitlines()
    matches=[i for i,x in enumerate(lines) if needle in x]
    assert len(matches)==1,(f,needle,matches)
    code(f,matches[0]+1,matches[0]+count)
def story(n,title,context):
    say(f'## Sprint {(n-1)//3+1} · HU{n:02d}: {title}\n\n{context}')

backlog=openpyxl.load_workbook(r'C:\Users\Usuario\Downloads\PRODUCT_BACKLOG_FINAL.xlsx',data_only=True)
rows=list(backlog['Backlog actualizado'].values)[3:12]
assert [int(r[1]) for r in rows]==[1,1,1,2,2,2,3,3,3]
assert all(f'HU{i:02d}' in rows[i-1][2] for i in range(1,10))
reference=Path(r'C:\Users\Usuario\Downloads\exposicion_sprints_1_2_3.md').read_text(encoding='utf-8-sig')
assert all(f'HU{i:02d}' in reference for i in range(1,10))

say('''# Guion de exposición técnica — Sprints 1, 2 y 3

**Uso:** las intervenciones están escritas para leerse en voz alta. Las indicaciones «En pantalla» ubican fragmentos literales del repositorio; no se leen como parte del discurso. Cada bloque es un rango continuo, sin código inventado ni omisiones insertadas. Los bloques breves pueden ser parte de una función mayor. Duración orientativa: 35–45 minutos, ampliable con la demo y preguntas.

**Fuentes:** PRODUCT_BACKLOG_FINAL.xlsx, «Backlog actualizado», filas 4–12; la extracción exposicion_sprints_1_2_3.md enviada por el usuario; y el código actual del repositorio SiderurgicaMovil. Los documentos se toman como referencias, no como instrucciones adicionales. Los criterios se contrastan con el código, no con el estado declarativo de la matriz.

### Apertura

Buenos días. Hoy voy a recorrer los avances de nuestros tres primeros sprints siguiendo la operación real de la aplicación. Primero veremos cómo se inicia y protege la sesión; después, cómo una cotización llega a pedido y venta; y finalmente, cómo gestionamos inventario, abastecimiento y captura de documentos.

En cada recorrido voy a partir de una acción del usuario y seguir la llamada hasta el componente que la resuelve. Después regresaré con la respuesta a la pantalla. Así podremos relacionar cada historia con una necesidad de negocio y con evidencia concreta en el código.

Nuestra aplicación es Android nativo con Kotlin y Jetpack Compose. Aquí el punto de entrada habitual es un evento de interfaz. El ViewModel coordina, el repositorio consume el ERP y Room conserva información en SQLite. Cuando lleguemos a Retrofit, veremos el contrato que usa la aplicación: el código del servidor ERP y sus transacciones no están en este repositorio. Por eso voy a separar lo que podemos demostrar en el cliente de lo que necesitamos validar en integración.
''')

story(1,'Base MVVM, Clean Code y Material Design 3','''Comienzo por HU01. La necesidad es tener una base consistente para que ventas, almacén y administración utilicen los mismos componentes y para que la lógica de una operación no quede mezclada con el dibujo de la pantalla. Voy a mostrarlo desde el arranque de Android, antes de entrar a una operación de negocio.''')
code('MainActivity.kt',21,31)
say('''Cuando Android crea la actividad, yo instalo el contenido Compose y entro a SiderurgicaScreen. No consulto una base de datos desde onCreate. Paso al componente que recibe el ViewModel y observa su estado.''')
code('MainActivity.kt',30,32)
say('''Aquí recojo el StateFlow respetando el ciclo de vida. Para entender qué está observando la pantalla, salto inmediatamente al ViewModel: allí se resuelve el repositorio y se construye el estado inicial con la sesión guardada.''')
code('ErpViewModel.kt',32,38)
say('''El repositorio y la base no se crean en cada recomposición. Los proporciona nuestra clase Application mediante propiedades lazy. Sigo esa referencia para ver cómo se instancian las dependencias.''')
code('SiderurgicaApp.kt',10,13)
say('''La base se llama siderurgica-mobile.db y Room administra su acceso. Esta es provisión manual de dependencias; no estamos mostrando un contenedor de inyección. Del lado remoto, el repositorio prepara Retrofit, añade el token Bearer si hay sesión y define un tiempo máximo de llamada.''')
code('data/ErpRepository.kt',65,72)
say('''Vuelvo ahora al dibujo de la interfaz. El componente raíz envuelve la navegación con SteelTheme; sigo ese nombre hasta el tema que aplica Material 3 y selecciona la paleta según el modo del sistema.''')
pick('MainActivity.kt','    SteelTheme {')
code('ui/Theme.kt',27,29)
say('''Con esto cierro HU01: tenemos una entrada Compose, estado observable, coordinación en ViewModel, acceso a datos mediante repositorio y base local Room. También tenemos tema claro y oscuro con Material 3. Esta historia establece una separación verificable de responsabilidades; no implica que ya tengamos todas las pruebas de accesibilidad o una arquitectura Clean completa certificada. Ahora voy a usar esta misma estructura para seguir la primera acción del usuario: ingresar.''')

story(2,'Autenticación, sesión y biometría','''HU02 resuelve el acceso seguro. Necesitamos identificar al usuario, recuperar su sesión y limitar sus operaciones por rol. El recorrido comienza cuando presiono Ingresar; el mismo botón cambia a registro cuando activo Crear cuenta.''')
pick('ui/HomeAndLogin.kt','                if (register) {',5)
say('''En el caso de ingreso llamamos a loginUnificado. Antes de ejecutar su cuerpo, la función entra al helper run. Sigo ese helper: evitamos lanzar otra operación mientras existe una carga activa y usamos viewModelScope para ejecutar la corrutina.''')
code('ErpViewModel.kt',46,50)
say('''Ahora vuelvo al cuerpo del login. Recibimos el identificador y la contraseña, quitamos espacios del correo y exigimos ambos campos. Si Firebase está configurado, validamos que el identificador tenga formato de correo.''')
code('ErpViewModel.kt',68,80)
say('''La validación invoca BusinessRules.validEmail, así que sigo esa llamada antes de avanzar. Aquí se verifica el formato; no estamos comprobando todavía la existencia de la cuenta.''')
code('domain/BusinessRules.kt',13,13)
say('''Superado ese control, extraigo el prefijo del correo como nombre de usuario ERP. Por ejemplo, admin@siderurgica.com se convierte en admin. En este flujo se usa la misma contraseña para el ERP y Firebase. La primera llamada de autenticación es al ERP.''')
code('ErpViewModel.kt',82,88)
code('data/ErpRepository.kt',74,80)
say('''Aquí enviamos LoginBody y comprobamos tres condiciones de la respuesta: éxito, usuario presente y token no vacío. Para ver adónde va la solicitud salto al contrato Retrofit.''')
code('data/ErpApi.kt',9,10)
code('data/ErpApi.kt',45,45)
say('''Este POST termina en el servidor ERP. Sus consultas de usuarios y comprobación de contraseña son externas a este código. Si recibimos una respuesta válida, also(vault::save) persiste la sesión; por eso sigo inmediatamente al almacén cifrado.''')
code('data/ErpRepository.kt',41,45)
say('''Convertimos la sesión a JSON y la ciframos con AES-GCM. La llave la obtenemos desde Android Keystore; no guardamos la contraseña ERP. Sigo key para mostrar dónde se crea o recupera esa llave.''')
code('data/ErpRepository.kt',31,40)
say('''Después del ERP regresamos al ViewModel y autenticamos Firebase. Solo entonces actualizamos el estado visible con la sesión y el correo Firebase.''')
code('ErpViewModel.kt',86,90)
say('''Si Firebase o el ERP fallan durante esa secuencia, deshacemos la sesión guardada y cerramos Firebase para no dejar una identidad parcial en la interfaz.''')
code('ErpViewModel.kt',91,100)
say('''En registro hay una condición de negocio que debemos dejar clara: primero verificamos que la cuenta ya exista en el ERP y después creamos su acceso Firebase. No estamos creando una cuenta empresarial desde el móvil.''')
code('ErpViewModel.kt',143,159)
say('''Para recuperar el acceso, el evento de pantalla invoca resetPassword. Sigo esa llamada: validamos el correo y solicitamos a Firebase el envío de recuperación. El cierre cancela la operación activa, borra el token local y cierra Firebase.''')
pick('ui/HomeAndLogin.kt','TextButton(onClick = { vm.resetPassword(identifier) }')
code('ErpViewModel.kt',168,173)
say('''La recuperación corresponde a Firebase; no hay una recuperación ERP en este contrato. Al reabrir, SessionVault.read descifra la sesión para restaurar el estado. Si falla la lectura, limpia el almacén y devuelve null.''')
code('data/ErpRepository.kt',46,56)
say('''Si el usuario habilitó biometría y el dispositivo la soporta, la pantalla bloqueada ofrece desbloquear. La confirmación del sensor cambia unlocked y deja continuar con la sesión previamente guardada.''')
pick('MainActivity.kt','                            val prompt = BiometricPrompt',4)
say('''Para las operaciones comerciales aplicamos además una barrera de rol en el repositorio. Sigo requireSales hasta su regla de dominio.''')
code('data/ErpRepository.kt',81,85)
code('domain/BusinessRules.kt',16,18)
say('''Cierro HU02 con ingreso, registro condicionado a una cuenta ERP, recuperación Firebase, cierre, sesión cifrada, roles y desbloqueo biométrico opcional. La interfaz controla acceso local, pero la autorización definitiva debe seguir existiendo en el servidor. La persistencia de sesión tampoco garantiza que un JWT siga vigente: un 401 debe tratarse como vencimiento. Con una identidad estable, paso a cómo mantenemos información útil cuando falla la conexión.''')

story(3,'Consulta offline y sincronización','''HU03 busca que una interrupción de internet no nos deje sin los últimos datos y que podamos preparar operaciones. Voy a seguir dos caminos: una lectura con caché y un borrador confirmado por el usuario.''')
pick('MainActivity.kt','vm::loadCache')
code('ErpViewModel.kt',275,281)
say('''El evento Consulta guardada llega a loadCache. Según el rol, pedimos listados con force=false. Para cotizaciones sigo la llamada al repositorio y luego a fetch, que es la implementación común de las lecturas.''')
code('data/ErpRepository.kt',90,90)
code('data/ErpRepository.kt',95,99)
say('''Al pedir datos locales, obtenemos el usuario propietario y consultamos LocalDao. La selección usa owner y kind para que no mezclemos cotizaciones de diferentes cuentas. Sigo la query hasta SQLite a través de Room.''')
code('data/LocalStore.kt',14,16)
code('data/LocalStore.kt',5,6)
say('''La respuesta de Room contiene registros JSON, que fetch convierte al modelo correspondiente. Regresamos al ViewModel, publicamos los listados y marcamos que estamos viendo datos guardados. Eso informa a la persona que el stock puede haber cambiado.

Si hacemos una consulta online, el mismo fetch llama primero al endpoint remoto y reemplaza la caché con la respuesta. Sigo esa rama y su operación transaccional local.''')
code('data/ErpRepository.kt',100,107)
code('data/LocalStore.kt',17,21)
say('''Aquí eliminamos el conjunto anterior de ese usuario y tipo y guardamos el nuevo dentro de una transacción Room. Ante una excepción no HTTP y no de cancelación, intentamos recuperar datos previos. Un 401 o un 403 se propaga; no lo ocultamos bajo una lectura local. Si no hay caché, también propagamos el error.

Para preparar una operación sin conexión, presiono Guardar como borrador. El formulario llama al mismo submit con draft=true y entrega el DTO al ViewModel.''')
pick('ui/SalesScreen.kt','OutlinedButton(onClick = { submit(true) }')
pick('ui/SalesScreen.kt','            fun submit(draft: Boolean)',6)
code('ErpViewModel.kt',282,286)
say('''En esta rama validamos igual que en línea, serializamos el payload y delegamos en saveDraft. Esta llamada no envía un POST.''')
code('data/ErpRepository.kt',151,154)
code('data/LocalStore.kt',8,12)
code('data/LocalStore.kt',22,24)
say('''Persistimos el borrador con UUID, usuario, tipo y fecha. Al volver a la pantalla aparece pendiente. Cuando el usuario confirma su envío, seguimos sendDraft: escogemos el endpoint por tipo y borramos el registro local solo después de que la llamada remota retorna sin error.''')
code('ErpViewModel.kt',297,306)
code('data/ErpRepository.kt',156,156)
code('data/LocalStore.kt',25,25)
say('''La parte periódica se agenda al inicializar el ViewModel con una restricción de red conectada. WorkManager administra cuándo puede ejecutarse; quince minutos es el intervalo configurado, no una promesa de ejecución exacta.''')
code('ErpViewModel.kt',41,43)
code('data/ErpRepository.kt',161,180)
say('''El worker consulta datos según rol y revisa stock cuando la consulta no usó caché. No consume ni envía los borradores: solo actualiza lecturas. Cierro esta historia con persistencia local, consulta offline y refresco periódico presentes. Queda pendiente el criterio de sincronizar escrituras automáticamente evitando duplicados. Una respuesta perdida después de un POST todavía puede causar un reenvío duplicado; la solución necesita idempotencia del servidor. Esta distinción nos acompaña ahora en el sprint comercial.''')

story(4,'Cotizaciones desde el móvil','''En HU04 el vendedor necesita preparar una propuesta, consultar sus importes y encontrarla por cliente, producto o estado. El punto de entrada es Confirmar cotización en CommercialForm. Antes de presionar, calculamos el resumen visible.''')
pick('ui/SalesScreen.kt','    val n = quantity.toIntOrNull()',3)
say('''La cantidad y el precio producen el subtotal. Si es cotización, llamamos a discount; salto a esa regla para ver la escala aplicada.''')
code('domain/BusinessRules.kt',15,15)
say('''Aplicamos 5 %, 7 % o 10 % desde los umbrales de diez, cincuenta y cien unidades. La pantalla calcula también IGV referencial al 18 %, pero presenta la cotización sin IGV y el pedido con IGV. El importe definitivo lo devuelve el ERP.

Al confirmar, submit prepara QuoteCreate con cliente, RUC, producto, cantidad, inventario seleccionado, precio y usuario. Sigo inmediatamente la llamada saveQuote.''')
pick('ui/SalesScreen.kt','                if(quote) vm.saveQuote')
code('ErpViewModel.kt',282,286)
say('''Primero exigimos rol comercial y validamos los campos. requireSales ya vimos que permite ADMIN o VENTAS; ahora sigo validate y su verificador de RUC.''')
code('ErpViewModel.kt',292,296)
code('domain/BusinessRules.kt',7,12)
say('''Aquí no basta con once dígitos: comprobamos el prefijo y el dígito de control del RUC. También rechazamos cantidades o precios no positivos y resultados no finitos. Con esto aprobado, regreso a createQuote en el repositorio.''')
code('data/ErpRepository.kt',128,134)
say('''Verificamos que el usuario del payload coincida con la sesión, llamamos a la API y exigimos una respuesta exitosa con cotización presente. Sigo el DTO y la ruta concreta de esa llamada.''')
code('data/ErpApi.kt',20,25)
code('data/ErpApi.kt',47,47)
say('''Los SerializedName adaptan los campos al contrato JSON. Retrofit ejecuta POST cotizaciones y devuelve QuoteCreated. La inserción en la base ERP ocurre fuera del repositorio móvil; aquí no tenemos su SQL.

Cuando retorna, el ViewModel publica el mensaje y vuelve a consultar cotizaciones. Sigo esa lectura hasta el GET y después el reemplazo local que ya recorrimos en HU03.''')
code('data/ErpRepository.kt',90,90)
code('data/ErpApi.kt',46,46)
pick('ui/SalesScreen.kt','    val quotes = state.quotes.filter')
say('''El listado actualizado entra al StateFlow y la pantalla filtra por texto y estado. Así cerramos el circuito: formulario, validación, POST, GET, caché y listado.

Cuando presiono Aprobar cotización, no hago una inserción directa de pedido desde la pantalla. Invoco approveQuote, que delega el cambio de estado y refresca cotizaciones y pedidos.''')
pick('ui/SalesScreen.kt','vm.approveQuote(quote.id)')
code('ErpViewModel.kt',308,311)
code('data/ErpRepository.kt',145,149)
code('data/ErpApi.kt',48,48)
say('''Luego, para una cotización APROBADA, la UI ofrece Crear pedido. Sigo esa acción hasta convertQuote y su endpoint específico.''')
pick('ui/SalesScreen.kt','vm.convertQuote(quote.id)')
code('ErpViewModel.kt',316,317)
code('data/ErpRepository.kt',139,139)
code('data/ErpApi.kt',49,49)
say('''Para rechazar, rejectQuote reutiliza quoteState con RECHAZADA. Concluyo HU04 con alta, consulta, búsqueda, filtros, descuento y transición comercial presentes. Trabajamos con un producto por operación, no con un carrito de líneas. También debo señalar que el README anterior habla de conversión automática al aprobar, mientras el código actual presenta aprobación y conversión separadas. En la demo sigo el código y dejo ese comportamiento del servidor como verificación de integración.''')

story(5,'Pedidos, aprobación y venta','''HU05 continúa el recorrido: la propuesta comercial se convierte en un compromiso de entrega y luego en una venta. Podemos partir de la conversión anterior o crear un pedido directo. Para este segundo camino el formulario llama createOrder.''')
pick('ui/SalesScreen.kt','                else vm.createOrder')
code('ErpViewModel.kt',287,291)
say('''Repetimos control de rol y validación de cliente, producto, cantidad, precio y RUC. Si no es borrador, seguimos al repositorio, que fija el identificador a cero y sustituye usuarioId por el propietario de la sesión.''')
code('data/ErpRepository.kt',135,138)
code('data/ErpApi.kt',51,51)
say('''El endpoint devuelve un Order. Al regresar, el ViewModel vuelve a consultar pedidos y publica Pedido registrado. La pantalla también muestra el último stock y bloquea la confirmación directa cuando la cantidad supera el disponible conocido; esa observación no reemplaza la validación concurrente del ERP.

Sobre un pedido pendiente, presiono Aprobar. Sigo el evento hasta orderState, el repositorio y el PATCH.''')
pick('ui/SalesScreen.kt','vm.orderState(order.id, "APROBADO")')
code('ErpViewModel.kt',312,315)
code('data/ErpRepository.kt',140,144)
code('data/ErpApi.kt',52,52)
say('''Aquí enviamos StateBody con APROBADO y comprobamos el código HTTP. Tras la respuesta recargamos pedidos, ventas e inventario. El móvil muestra el resultado del servidor; no ejecuta la reserva o el descuento mediante una query local. ANULADO y RECHAZADO utilizan la misma cadena con distinto estado.

Cuando el pedido queda APROBADO, la interfaz ofrece Generar venta. Sigo esa acción sin detenerme en otro archivo ajeno al flujo.''')
pick('ui/SalesScreen.kt','vm.sale(order.id)')
code('ErpViewModel.kt',318,318)
code('data/ErpRepository.kt',150,150)
code('data/ErpApi.kt',54,54)
say('''La API devuelve Sale con cliente, producto, cantidad e importes. Después de crear la venta recargamos pedidos y ventas para que la pantalla refleje el resultado. En esta función no recargamos inventario; la recarga de inventario está en orderState y en el refresco general.''')
code('data/ErpApi.kt',34,39)
say('''Cierro HU05 con registro, cambios de estado y generación de venta visibles en el código. No afirmo que las transacciones o la actualización de stock del ERP estén verificadas: faltan sus fuentes y una prueba integrada. Tampoco tenemos reintentos seguros de escritura; OkHttp está configurado sin retryOnConnectionFailure. El README describía venta automática al aprobar y la UI actual ofrece una acción posterior, por lo que esa secuencia debe contrastarse con el servidor real. Ahora tomo la venta recibida para generar su constancia.''')

story(6,'PDF, vista previa, compartir y Storage','''HU06 busca que el usuario tenga una constancia digital de la venta. El recorrido parte de la tarjeta de venta y su botón Vista previa PDF o Compartir constancia. Aquí hay una particularidad de arquitectura: la pantalla llama directamente al servicio SalePdf; no pasa por un método del ViewModel.''')
pick('ui/SalesScreen.kt','OutlinedButton(onClick = { pdfAction(true) }',2)
code('ui/SalesScreen.kt',110,118)
say('''Ambas acciones entran a pdfAction y crean el archivo. Sigo SalePdf.create: ubicamos el PDF en el directorio de caché y abrimos un documento de una página.''')
code('SalePdf.kt',20,27)
say('''Con el modelo Sale ya obtenido del ERP dibujamos cliente, RUC, producto, cantidad e importes. El PDF no vuelve a calcular el total; representa los valores de la venta.''')
code('SalePdf.kt',35,47)
say('''Al finalizar escribimos el documento al archivo y lo devolvemos al llamador. Si la acción era vista previa, pdfAction entrega una URI FileProvider a ACTION_VIEW. Para compartir sigo share, que usa ACTION_SEND y permite leer temporalmente la URI.''')
code('SalePdf.kt',49,55)
say('''Así dejamos que Android muestre sus destinos nativos de compartición. Para almacenar el PDF en Firebase, presiono Subir a Firebase. El evento lanza una corrutina y llama upload con el archivo recién creado.''')
pick('ui/SalesScreen.kt','                try { SalePdf.upload')
code('SalePdf.kt',56,63)
say('''Tomamos el uid de Firebase, definimos application/pdf y subimos a usuarios/uid/ventas/id.pdf. La autorización de ese destino se ve en las reglas incluidas; sigo inmediatamente ese control.''')
code('firebase.storage.rules',4,10)
say('''Solo el usuario autenticado correspondiente puede leer o escribir en esa ruta, y la escritura exige tamaño menor a diez megabytes y tipo PDF. La pantalla comunica el resultado y sale del estado Subiendo en finally.

Cierro HU06 con generación, apertura mediante visor externo, compartir y subida de una constancia informativa. El criterio de boleta o factura tributaria todavía no queda satisfecho: el propio documento dice que no es comprobante de pago tributario. No hay numeración fiscal ni integración de emisión en este código. Con las operaciones comerciales recorridas, paso al inventario que las sostiene.''')

story(7,'Inventario, movimientos y alertas de stock','''HU07 atiende una pregunta operativa: qué podemos vender y cuándo necesitamos reponer. Entramos al módulo Stock; el refresco pide inventario al repositorio y la pantalla usa el disponible o, si falta, el total.''')
pick('MainActivity.kt','"stock" -> InventoryScreen')
code('ErpViewModel.kt',231,233)
code('data/ErpRepository.kt',93,93)
code('data/ErpApi.kt',55,55)
say('''Esta lectura pasa por fetch: al retornar del ERP reemplaza la caché por usuario. Con los datos ya publicados, la tarjeta invoca stockLevel y la regla clasifica normal, bajo, agotado o sin dato.''')
pick('ui/InventoryAndPurchases.kt','StatusBadge(BusinessRules.stockLevel')
code('domain/BusinessRules.kt',19,24)
say('''El umbral viene del stock mínimo del producto; usamos diez como respaldo cuando falta. Para modificar existencias, el botón Registrar movimiento abre un diálogo. Allí selecciono entrada, salida o ajuste y confirmo.''')
pick('ui/InventoryAndPurchases.kt','TextButton(onClick = { submitted = true; vm.saveMovement(item, kind, quantity.toIntOrNull() ?: 0, reason) }')
code('ErpViewModel.kt',351,362)
say('''Solo ALMACEN o ADMIN pueden operar. Validamos el tipo, cantidad y motivo y comparamos una salida con el último disponible conocido. Si guardo localmente, inserto StockMovement; sigo esa entidad y el DAO.''')
code('data/Operations.kt',9,10)
pick('data/Operations.kt','@Insert suspend fun addMovement')
say('''Esta inserción deja una solicitud pendiente en SQLite; no altera el stock ERP. Si confirmo en línea, seguimos changeStock. Un ajuste establece un valor final; una entrada o salida expresa una cantidad de movimiento.''')
code('data/ErpRepository.kt',118,126)
say('''Para AJUSTE consultamos primero el producto remoto y evitamos bajar el total por debajo del stock reservado conocido. Después hacemos PUT. Para entrada o salida hacemos PATCH con StockBody; sigo ambos contratos.''')
code('data/ErpApi.kt',56,58)
code('data/ErpApi.kt',70,72)
say('''Comprobamos success y al retornar volvemos a consultar inventario. El motivo queda en el borrador local, pero no existe un campo de motivo en StockBody, de modo que no se envía al ERP.

Para enviar una solicitud local, seguimos sendMovement: ejecutamos el mismo servicio, eliminamos el borrador tras la respuesta y actualizamos el listado.''')
code('ErpViewModel.kt',398,402)
say('''La otra rama de HU07 es la alerta. El CacheSyncWorker que mostramos en HU03 obtiene stock y, si no utilizó caché, llama StockNotifications.check. Sigo esa llamada: guardamos el nivel anterior y notificamos cuando el producto pasa a Bajo o Agotado.''')
pick('data/ErpRepository.kt','if (!repo.usedCache) pe.edu.unmsm.siderurgica.StockNotifications.check')
code('domain/StockNotifications.kt',27,37)
say('''check invoca show, que comprueba el permiso de notificaciones en Android 13 o posterior, crea el canal y publica la notificación local.''')
code('domain/StockNotifications.kt',18,26)
say('''También existe un receptor FCM. Si llega un mensaje, verifica la sesión y el rol y pide abrir la app para consultar el inventario vigente.''')
code('domain/StockNotifications.kt',41)
say('''Concluyo HU07 con consulta, clasificación visual, movimientos, solicitudes locales y alertas periódicas en el dispositivo. La recepción FCM está implementada, pero el token solo se conserva localmente: no hay endpoint para registrarlo en el ERP ni emisor de alertas del servidor. Por tanto, el circuito push completo queda pendiente. Desde esta necesidad de reposición entramos a proveedores y órdenes de compra.''')

story(8,'Proveedores, órdenes de compra y recepción','''HU08 resuelve el abastecimiento: registrar a quién compramos, controlar el importe preparado y reconocer cuándo la mercadería realmente ingresa. Comienzo creando un proveedor; después uso ese proveedor para una orden y finalmente registro su recepción.''')
pick('ui/InventoryAndPurchases.kt','TextButton(onClick = { submitted = true; vm.saveSupplier(name, ruc, contact, phone, "") }')
code('ErpViewModel.kt',321,329)
say('''Aquí recibimos razón social, RUC y contacto. Aplicamos la validación de RUC que vimos en HU04 y consultamos OperationsDao para detectar duplicados locales; también revisamos la lista cargada.''')
pick('data/Operations.kt','suspend fun suppliers(owner: String)')
say('''Si elijo borrador, addSupplier persiste la entidad. Su índice único por owner y ruc añade una barrera local de integridad.''')
code('data/Operations.kt',5,6)
pick('data/Operations.kt','@Insert suspend fun addSupplier')
say('''Si registro en línea, además de exigir conexión consultamos proveedores remotos, verificamos el RUC y llamamos createSupplier. Sigo esa llamada hasta el POST.''')
code('ErpViewModel.kt',330,336)
code('data/ErpRepository.kt',114,114)
code('data/ErpApi.kt',60,60)
say('''La respuesta retorna un proveedor ERP y después refrescamos el listado. Ya con un proveedor remoto seleccionado, presiono Registrar orden ERP. El diálogo entrega producto, cantidad, precio y presupuesto a savePurchase.''')
pick('ui/InventoryAndPurchases.kt','TextButton(onClick = { submitted = true; vm.savePurchase(supplier, product, n, p, b) }')
code('ErpViewModel.kt',337,350)
say('''Validamos el propietario del proveedor y que cantidad por precio no exceda el límite ingresado. Una orden online requiere serverId. Antes del POST redondeamos el importe con amount; sigo ese helper porque forma parte del payload enviado.''')
code('ErpViewModel.kt',414,414)
code('data/ErpRepository.kt',115,115)
code('data/ErpApi.kt',76,76)
code('data/ErpApi.kt',62,62)
say('''Enviamos la cantidad como texto porque ese es el contrato de PurchaseOrderBody. El precio y el presupuesto no son campos de ese DTO; mandamos el total calculado. La UI consulta el presupuesto ERP y lo muestra, pero el control transaccional de presupuesto debe confirmarse en el servidor.

Al volver recargamos órdenes y presupuesto. Si preparé la compra offline, la rama asDraft inserta Purchase en Room y no reserva fondos. Sigo la persistencia correspondiente.''')
code('data/Operations.kt',7,8)
pick('data/Operations.kt','@Insert suspend fun addPurchase')
say('''El envío del borrador es explícito: buscamos el proveedor ERP por identificador o RUC, mandamos la orden y solo después eliminamos la compra local.''')
code('ErpViewModel.kt',390,397)
say('''Antes de recibir, el usuario aprueba la orden. La acción salta a purchaseState, delega al repositorio y usa PATCH para cambiar su estado.''')
pick('ui/InventoryAndPurchases.kt','vm.purchaseState(order, "Aprobada")')
code('ErpViewModel.kt',403,405)
code('data/ErpRepository.kt',116,116)
code('data/ErpApi.kt',63,63)
say('''La recepción se habilita para una orden aprobada o en tránsito que todavía no tenga recepción cargada. Cuando la mercadería llega, presiono Recibir y actualizar stock. Sigo el método invocado.''')
pick('ui/InventoryAndPurchases.kt','vm.receivePurchase(order, n, p)')
code('ErpViewModel.kt',406,413)
say('''Controlamos cantidades y precios positivos, comparamos importes redondeados y consultamos recepciones existentes para evitar registrar otra sobre la misma orden. Esa consulta llega a receipts y después al GET compra.''')
code('data/ErpRepository.kt',112,112)
code('data/ErpApi.kt',64,64)
say('''Si pasa la validación, construimos ReceiptBody con el producto y, cuando lo encontramos, su inventarioId. Ahora sigo receivePurchase al POST y a la estructura que retorna.''')
code('data/ErpRepository.kt',117,117)
code('data/ErpApi.kt',65,65)
code('data/ErpApi.kt',79,79)
say('''El contrato devuelve compra e inventario junto con indicadores de creación de producto y stock anterior. El ViewModel no usa esos valores para hacer un ajuste local: vuelve a consultar recepciones, órdenes e inventario y publica Compra recibida.

Cierro HU08 con proveedores validados, órdenes, borradores manuales, cambio de estado y recepción conectada a la recarga de inventario. La consulta de duplicados en cliente ayuda, pero no garantiza exclusión bajo concurrencia; eso necesita restricciones y transacciones del servidor. La cola local no se envía automáticamente. Ahora voy a capturar el documento que respalda una operación.''')

story(9,'Captura de documentos, OCR y códigos','''HU09 reduce la digitación de facturas y permite leer códigos QR o de barras. El recorrido comienza en Fotografiar documento o código. Creamos un destino local y pedimos a Android capturar una imagen; también podemos elegirla desde la galería.''')
pick('ui/Documents.kt','                    val file = File(context.cacheDir',4)
pick('ui/Documents.kt','    val camera = rememberLauncherForActivityResult',2)
say('''Cuando TakePicture retorna éxito, ejecutamos process con la URI. Dentro de una corrutina llamamos DocumentScanner.scan. Sigo el servicio antes de volver al formulario.''')
code('ui/Documents.kt',38,44)
code('data/DocumentScanner.kt',20,28)
say('''Leemos las dimensiones, rechazamos imágenes inválidas y reducimos el tamaño de decodificación cuando la dimensión mayor supera el límite usado para el muestreo. Después leemos EXIF para aplicar giro y reflejo; así el reconocimiento y la imagen guardada conservan la misma orientación.''')
code('data/DocumentScanner.kt',29,39)
say('''Dentro de scan creamos el reconocedor de texto y el de códigos ML Kit. Procesamos el bitmap, obtenemos el texto y los valores de los códigos y comprimimos la imagen a JPEG con calidad ochenta.''')
code('data/DocumentScanner.kt',40,46)
say('''La última línea invoca InvoiceParser.parse; sigo esa llamada antes de explicar la vuelta a la pantalla. El parser busca RUC, número, fecha e importe total dentro del texto reconocido.''')
code('domain/BusinessRules.kt',27,37)
say('''La selección de RUC vuelve a usar validRuc. Lo que hemos mostrado es OCR de ML Kit seguido por extracción con expresiones regulares; no es un modelo generativo que interprete toda la factura. En scan, finally libera el bitmap y cierra los dos clientes.

Regresamos con ScanResult. process propone los campos y normaliza separadores de fecha, pero el usuario conserva el control: puede revisar la imagen y corregir antes de confirmar. El botón Confirmar factura entrega los datos revisados a saveDocument.''')
pick('ui/Documents.kt','vm.saveDocument(imagePath, ruc, number, date')
code('ErpViewModel.kt',363,370)
say('''Validamos otra vez RUC y número, exigimos una fecha real con formato estricto, monto positivo y archivo existente. Después delegamos en addDocument; sigo esa inserción y su entidad Room.''')
code('data/Operations.kt',11,12)
pick('data/Operations.kt','@Insert suspend fun addDocument')
say('''Guardamos la ruta de imagen y los campos en SQLite, asociados al usuario ERP. Después localData vuelve a leerlos y la pantalla muestra el documento. Si solo encontramos un código, saveCode usa la misma entidad, sin inventar RUC ni monto de una factura.''')
code('ErpViewModel.kt',371,375)
say('''La subida a Firebase es una acción independiente del guardado. El botón Subir imagen a Firebase invoca uploadDocument. Sigo esa llamada desde la comprobación de propietario hasta la escritura remota.''')
pick('ui/Documents.kt','vm.uploadDocument(document)')
code('ErpViewModel.kt',376,383)
say('''Usamos el uid Firebase para la ruta usuarios/uid/documentos/id.jpg y declaramos image/jpeg. Solo después de que putFile finaliza invocamos markUploaded. Sigo ese cambio para ver cómo se registra el resultado en la base local.''')
pick('data/Operations.kt','suspend fun markUploaded')
say('''La query actualiza storagePath únicamente para el documento y su propietario. Después recargamos los datos y el estado visual pasa de En dispositivo a En Firebase.

Para cerrar este recorrido debo mostrar una dependencia pendiente: las reglas Storage que incluye el repositorio habilitan PDFs en ventas, pero no declaran la ruta documentos. No basta con tener putFile escrito para considerar autorizada esa subida.''')
code('firebase.storage.rules',1)
say('''Cierro HU09 con captura por cámara o galería, OCR local, lectura de códigos, compresión, revisión humana y persistencia Room. La llamada de subida y el registro local de su resultado existen, pero con estas reglas la subida de JPEG a documentos no queda habilitada. Esa configuración y su prueba son pendientes concretos de integración; no debemos presentar el almacenamiento remoto de imágenes como una capacidad ya validada.''')

say('''## Cierre de la exposición

Con este recorrido completamos las nueve historias de los tres primeros sprints. En el primero establecimos arquitectura, identidad y continuidad de consulta; en el segundo conectamos cotizaciones, pedidos, ventas y constancias; en el tercero extendimos la operación a stock, abastecimiento y captura documental.

Lo que mostramos tiene una trazabilidad concreta: el evento llega a la coordinación, la coordinación valida y delega, el dato se consulta o persiste y el resultado vuelve al estado observado por Compose. Para el equipo técnico, esa separación permite ubicar tanto una regla como una dependencia. Para los stakeholders, permite distinguir preparar una operación de confirmarla en el ERP y conocer qué resultado obtiene el usuario.

Los siguientes cierres necesarios están identificados: idempotencia para sincronizar escrituras automáticamente, validación integrada de las transacciones ERP y de la secuencia aprobación–conversión–venta, emisión tributaria si exigimos boleta o factura, registro y emisión de FCM en servidor, y reglas Storage que permitan imágenes documentales. También necesitamos ejecución y pruebas integradas en Android; este guion verifica fuentes y llamadas, no certifica una demo realizada.

Con eso termino la presentación. Podemos pasar a ejecutar estos recorridos con cuentas de prueba y contrastar las respuestas del ERP, la persistencia local y los resultados de Firebase.

### Notas para conducir la demo

Antes de comenzar, preparar cuentas ERP con los roles usados y, cuando corresponda, accesos Firebase compatibles con el flujo unificado; configurar la URL ERP y conectividad. Para probar offline, cargar primero datos en línea y después cortar conexión. Una solicitud guardada no debe confundirse con stock modificado ni un borrador con una operación ERP confirmada.

Mostrar el PDF como constancia informativa. Mostrar FCM como receptor disponible y las alertas periódicas como notificaciones locales. Para HU09, demostrar captura y guardado local; presentar la subida como pendiente si se utilizan las reglas incluidas. Ante una respuesta perdida de escritura, comprobar el listado antes de reenviar. Estas indicaciones mantienen la demo alineada con el alcance descrito.
''')

for p,a,b,s in verified:
    assert s=='\n'.join(p.read_text(encoding='utf-8-sig').splitlines()[a-1:b])
out=root/'outputs'/'guion_walkthrough_sprints_1_2_3.md'
out.parent.mkdir(exist_ok=True)
out.write_text('\n'.join(parts),encoding='utf-8')
print(f'{out}\n{len(verified)} fragmentos literales verificados; {len(out.read_text(encoding="utf-8").split())} palabras totales')
